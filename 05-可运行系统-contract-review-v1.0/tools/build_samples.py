# -*- coding: utf-8 -*-
"""示例合同集生成器（构建期工具，运行一次即可，运行系统本身不需要本脚本）

依赖：python-docx / reportlab / pillow / pypdf —— 仅构建期使用。
产物：
    samples/01_软件采购合同_风险用例.docx        DOCX 风险用例
    samples/02_服务器设备采购合同_标准用例.pdf    PDF 文本型 标准用例
    samples/03_大额销售合同_违约责任不对等.docx    DOCX 风险用例（违约责任不对称）
    samples/04_技术服务合同_标准用例.pdf          PDF 文本型 标准用例
    samples/05_劳动合同_风险用例.docx             DOCX 劳动用工风险
    samples/06_软件采购合同_扫描件.png            图片扫描件（走 Windows 内置 OCR）
    samples/07_设备采购合同_无文本层.pdf          图片型 PDF（无文本层 → 解析受阻）
    samples/08_保密协议_加密用例.pdf              加密 PDF（→ 解析受阻）

用法：
    python tools/build_samples.py
"""
from __future__ import annotations

import io
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tools"))

from contracts import C01, C02, C03, C04, C05, C06_LINES          # noqa: E402

OUT = os.path.join(ROOT, "samples")
os.makedirs(OUT, exist_ok=True)

C07_LINES = [
    "《备品备件采购合同》",
    "合同编号：CR-2026-0918-007",
    "甲方（需方）：华北智造集团有限公司",
    "乙方（供方）：中恒物资贸易有限公司",
    "第一条 标的物与金额",
    "采购内容：液压泵、轴承、密封件等备品备件一批，具体清单见附件一。",
    "合同金额：人民币 96,000 元，含 13% 增值税。",
    "第二条 交付",
    "乙方应于合同生效之日起 30 日内将货物送达甲方指定仓库，运输费用由乙方承担。",
    "第三条 合同生效",
    "本合同自双方签字并加盖公章之日起生效，一式肆份，双方各执贰份。",
]

C08 = [
    "《技术保密协议》",
    "合同编号：CR-2026-0918-008",
    "甲方：华北智造集团有限公司",
    "乙方：启元数据科技有限公司",
    "第一条 保密信息",
    "本协议所称保密信息包括双方在技术交流与项目合作过程中披露的技术方案、源代码、数据模型、经营数据与客户信息。",
    "第二条 保密义务",
    "双方对保密信息承担保密义务，未经披露方书面同意不得向第三方披露、复制或用于本协议约定目的之外的用途。",
    "第三条 保密期限",
    "保密期限自本协议生效之日起至保密信息公开之日或本协议终止后五年届满之日，以较晚者为准。",
    "第四条 违约责任",
    "任一方违反本协议约定披露保密信息的，应向守约方赔偿因此造成的实际损失。",
    "第五条 协议生效",
    "本协议自双方签字并加盖公章之日起生效，一式肆份，双方各执贰份。",
]


def log(msg):
    print(msg)


# ---------------------------------------------------------------- DOCX
def build_docx(path, lines):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if i == 0:
            run.bold = True
            run.font.size = Pt(15)
            p.alignment = 1
        elif line.startswith(("第", "合同编号", "甲方", "乙方", "统一社会信用代码", "住所", "法定代表人",
                              "身份证号码", "住址")):
            if line.startswith("第"):
                run.bold = True
    doc.save(path)
    log("  [docx] %s" % os.path.basename(path))


# ---------------------------------------------------------------- PDF（文本型）
_FONT_CACHE = {"name": None}


def _register_cjk_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if _FONT_CACHE["name"]:
        return _FONT_CACHE["name"]
    candidates = [
        ("MSYH", r"C:\Windows\Fonts\msyh.ttc", 0),
        ("MSYH", r"C:\Windows\Fonts\msyh.ttf", None),
        ("SIMSUN", r"C:\Windows\Fonts\simsun.ttc", 0),
        ("SIMHEI", r"C:\Windows\Fonts\simhei.ttf", None),
    ]
    for name, path, idx in candidates:
        if not os.path.exists(path):
            continue
        try:
            if idx is None:
                pdfmetrics.registerFont(TTFont(name, path))
            else:
                pdfmetrics.registerFont(TTFont(name, path, subfontIndex=idx))
            _FONT_CACHE["name"] = name
            log("  [font] 使用 TrueType 字体 %s（含 ToUnicode 映射）" % os.path.basename(path))
            return name
        except Exception as exc:                                   # noqa: BLE001
            log("  [font] %s 注册失败：%s" % (os.path.basename(path), exc))
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    _FONT_CACHE["name"] = "STSong-Light"
    log("  [font] 回退到 CID 字体 STSong-Light")
    return "STSong-Light"


def build_pdf(path, lines, title_center=True):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    font = _register_cjk_font()
    c = canvas.Canvas(path, pagesize=A4)
    W, H = A4
    left, right = 20 * mm, W - 20 * mm
    y = H - 22 * mm
    leading = 7.0 * mm
    max_chars = 42

    def wrap(text):
        out, cur = [], ""
        for ch in text:
            cur += ch
            if len(cur) >= max_chars:
                out.append(cur)
                cur = ""
        if cur:
            out.append(cur)
        return out

    for i, line in enumerate(lines):
        c.setFont(font, 15 if i == 0 else 10.5)
        blocks = wrap(line)
        for j, blk in enumerate(blocks):
            if y < 20 * mm:
                c.showPage()
                c.setFont(font, 10.5)
                y = H - 22 * mm
            if i == 0 and title_center and j == 0:
                c.drawCentredString(W / 2.0, y, blk)
            else:
                c.drawString(left, y, blk)
            y -= leading
        if i == 0:
            y -= 3 * mm
    c.save()
    log("  [pdf ] %s（%d 行）" % (os.path.basename(path), len(lines)))


# ---------------------------------------------------------------- 扫描件 PNG
def build_scan_png(path, lines, width=1600):
    from PIL import Image, ImageDraw, ImageFont

    font_path = r"C:\Windows\Fonts\msyh.ttc"
    if not os.path.exists(font_path):
        font_path = r"C:\Windows\Fonts\simsun.ttc"
    title_font = ImageFont.truetype(font_path, 46, index=0)
    body_font = ImageFont.truetype(font_path, 33, index=0)

    margin_x, margin_y, leading = 115, 100, 60
    body_w = width - margin_x * 2
    wrapped = []
    for i, line in enumerate(lines):
        f = title_font if i == 0 else body_font
        cur = ""
        for ch in line:
            cur += ch
            if f.getlength(cur) > body_w:
                wrapped.append((cur[:-1], f))
                cur = ch
        wrapped.append((cur, f))
    height = margin_y * 2 + leading * len(wrapped) + 60

    img = Image.new("RGB", (width, height), (252, 251, 248))
    d = ImageDraw.Draw(img)
    y = margin_y
    for i, (text, f) in enumerate(wrapped):
        if i == 0:
            w = f.getlength(text)
            d.text(((width - w) / 2, y), text, font=f, fill=(24, 24, 26))
            y += leading + 20
            continue
        head = text.startswith("第")
        if head:
            y += 10
        d.text((margin_x, y), text, font=f,
               fill=(14, 14, 16) if head else (32, 32, 36))
        y += leading
    # 扫描边缘轻微压暗
    for k in range(10):
        v = 246 - k * 3
        d.rectangle([k, k, width - 1 - k, height - 1 - k], outline=(v, v, v))
    img.save(path, "PNG", optimize=True)
    log("  [png ] %s（%dx%d）" % (os.path.basename(path), width, height))
    return path


# ---------------------------------------------------------------- 图片型 PDF（无文本层）
def build_image_pdf(path, png_path):
    from PIL import Image
    img = Image.open(png_path)
    if img.mode != "RGB":
        img = img.convert("RGB")
    img.save(path, "PDF", resolution=150.0)
    log("  [pdf ] %s（图片型 PDF，无文本层）" % os.path.basename(path))


# ---------------------------------------------------------------- 加密 PDF
def build_encrypted_pdf(path, src_pdf):
    from pypdf import PdfReader, PdfWriter
    reader = PdfReader(src_pdf)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password="contract2026", owner_password="owner2026",
                   algorithm="AES-128")
    with open(path, "wb") as f:
        writer.write(f)
    log("  [pdf ] %s（AES-128 加密，口令 contract2026）" % os.path.basename(path))


def main():
    log("生成示例合同集 -> %s" % OUT)
    build_docx(os.path.join(OUT, "01_软件采购合同_风险用例.docx"), C01)
    build_pdf(os.path.join(OUT, "02_服务器设备采购合同_标准用例.pdf"), C02)
    build_docx(os.path.join(OUT, "03_大额销售合同_违约责任不对等.docx"), C03)
    build_pdf(os.path.join(OUT, "04_技术服务合同_标准用例.pdf"), C04)
    build_docx(os.path.join(OUT, "05_劳动合同_风险用例.docx"), C05)

    png = build_scan_png(os.path.join(OUT, "06_软件采购合同_扫描件.png"), C06_LINES)
    png7 = build_scan_png(os.path.join(OUT, "_tmp07.png"), C07_LINES)
    build_image_pdf(os.path.join(OUT, "07_设备采购合同_无文本层.pdf"), png7)
    try:
        os.remove(png7)
    except OSError:
        pass

    tmp8 = os.path.join(OUT, "_tmp08.pdf")
    build_pdf(tmp8, C08)
    build_encrypted_pdf(os.path.join(OUT, "08_保密协议_加密用例.pdf"), tmp8)
    try:
        os.remove(tmp8)
    except OSError:
        pass

    log("完成。共 %d 个示例文件。" % len([f for f in os.listdir(OUT)
                                      if not f.startswith("_")]))


if __name__ == "__main__":
    main()
